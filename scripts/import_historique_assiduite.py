"""Import de l'historique d'assiduité depuis les « Livres des chantiers » Excel.

Reprend les années 6022-6023 à 6024-6025 (de l'allumage des feux du 21/12/2022
jusqu'à juin 2025) : tenues, pointages des membres (P/E/A), maçons passants.
L'année 6025-6026 est déjà saisie sur le portail : les onglets correspondants
sont ignorés.

Le script passe par les modèles de l'application et fonctionne donc sur la base
pointée par DATABASE_URL (SQLite ou PostgreSQL). Il est idempotent : une tenue,
un pointage ou une visite déjà présents ne sont jamais recréés.

Usage :
    python scripts/import_historique_assiduite.py MEMBRES.xlsx PASSANTS.xlsx           # simulation
    python scripts/import_historique_assiduite.py MEMBRES.xlsx PASSANTS.xlsx --apply   # écriture
"""
import argparse
import asyncio
import re
import sys
import unicodedata
from collections import Counter, defaultdict
from datetime import date, datetime
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))

import openpyxl
from sqlalchemy import select
from sqlalchemy.ext.asyncio import AsyncSession, async_sessionmaker, create_async_engine

from app.config import get_settings
from app.models.documents import Document
from app.models.identity import MasonicGrade, Member, MemberStatus
from app.models.lodge import MasonicYear
from app.models.meetings import (
    Attendance, AttendanceStatus, Meeting, MeetingGrade, MeetingType,
    MeetingVisitor, Visitor, VisitorStatus,
)
from app.models.reports import MeetingReport, ReportStatus
import app.models  # noqa: F401 — enregistre toutes les tables (clés étrangères)


# ── Paramètres de la reprise (validés avec le V∴M∴) ─────────────────────────

SHEETS = ["2022-2023", "2023-2024", "2024-2025"]
ALLUMAGE_DES_FEUX = date(2022, 12, 21)

# Erreurs de date dans les fichiers, corrigées d'après les tracés
MEMBRES_DATE_FIXES = {date(2023, 2, 15): date(2023, 2, 13)}    # tracé 004 : « Tenue du 13 février 2023 »
PASSANTS_DATE_FIXES = {date(2023, 11, 18): date(2023, 11, 15)}  # tracé 017 : 15 novembre 2023

# Nom dans l'Excel → (nom, prénom) sur le portail
MEMBER_ALIASES = {
    ("ANOUGNA ONANA", "Alphonse"): ("AGNOUNA ONANA", "Alphonse"),
    ("CANEL-GUERARD", "Aline"): ("CANEL", "Aline"),
    ("TRIEU", "Jean-Baptiste"): ("LOUIS-TRIEU", "Jean-Baptiste"),
    ("DANGOUMAU", "Tristan Jacques"): ("Dangoumau", "Tristan"),
    ("KAISER", "Nicolas"): ("Kaiser", "Nicolas"),
}

# Corrections de fiches membres existantes
RENAMES = {("Kaizer", "Nicolas"): ("Kaiser", "Nicolas")}
# Date d'entrée générique 22/12/2022 → jour de l'allumage des feux
GENERIC_START_FIX = (date(2022, 12, 22), ALLUMAGE_DES_FEUX)
# (nom, prénom) → champs à corriger
MEMBER_FIXES = {
    ("CANEL", "Aline"): {"membership_start_date": date(2025, 6, 21)},  # affiliation
    # Réception le 17/04/2024 (tracé 027) ; le 04/04/2024 était une tenue en Chambre du Milieu
    ("LEROY", "Christine"): {"initiation_date": date(2024, 4, 17), "membership_start_date": date(2024, 4, 17)},
    ("MANGENOT", "Gerard"): {"initiation_date": date(2024, 4, 17), "membership_start_date": date(2024, 4, 17)},
}

# Membres partis, absents du portail : (nom, prénom) → (civilité, date de démission)
DEPARTED = {
    ("CLERC", "Florence"): ("S", date(2022, 12, 31)),
    ("CHANSON", "Claire"): ("S", date(2023, 12, 31)),
    ("ROCHE", "Dominique"): ("F", date(2023, 12, 31)),
    ("LUZERNE", "Christian"): ("F", date(2024, 10, 1)),
    ("BUR", "Patrick"): ("F", date(2024, 12, 1)),
    ("SEYLER", "Synthia"): ("S", date(2025, 9, 1)),
}

# Tenues absentes des livres des chantiers mais attestées par un tracé (sans pointage)
EXTRA_MEETINGS = [(date(2025, 5, 21), MeetingGrade.COMPAGNON, 49)]

# Tenues de table (suffixe « Table » des tracés)
TABLE_MEETINGS = {(date(2024, 6, 19), MeetingGrade.APPRENTI), (date(2025, 6, 21), MeetingGrade.APPRENTI)}

GRADE_FROM_EXCEL = {"APP": MeetingGrade.APPRENTI, "COMP": MeetingGrade.COMPAGNON, "M": MeetingGrade.MAITRE}
STATUS_FROM_EXCEL = {"P": AttendanceStatus.PRESENT, "E": AttendanceStatus.EXCUSED, "A": AttendanceStatus.ABSENT}
VISITOR_GRADE = {"M": "Maître", "VM": "Maître", "C": "Compagnon", "A": "Apprenti", "APP": "Apprenti"}

JOURS = ["lundi", "mardi", "mercredi", "jeudi", "vendredi", "samedi", "dimanche"]
MOIS = ["janvier", "février", "mars", "avril", "mai", "juin", "juillet", "août",
        "septembre", "octobre", "novembre", "décembre"]
MOIS_ABR = {"jan": 1, "fev": 2, "mar": 3, "avr": 4, "mai": 5, "jun": 6, "juin": 6,
            "jul": 7, "juil": 7, "aou": 8, "sep": 9, "oct": 10, "nov": 11, "dec": 12}


# ── Helpers ──────────────────────────────────────────────────────────────────

def ascii_key(s) -> str:
    s = unicodedata.normalize("NFKD", str(s or "")).encode("ascii", "ignore").decode().upper()
    return re.sub(r"[^A-Z]", "", s)


def person_key(last, first) -> tuple[str, str]:
    return ascii_key(last), ascii_key(first)


def split_excel_name(full: str) -> tuple[str, str]:
    """« DEL VECCHIO Jean-Pierre » → (« DEL VECCHIO », « Jean-Pierre »)."""
    tokens = str(full).split()
    last = [t for t in tokens if t.isupper()]
    first = [t for t in tokens if not t.isupper()]
    return " ".join(last), " ".join(first)


def label_key(s) -> str:
    """Clé de rapprochement des noms de loges / orients (casse, accents, articles)."""
    s = unicodedata.normalize("NFKD", str(s or "")).encode("ascii", "ignore").decode().lower()
    s = re.sub(r"\bst\b\.?", "saint", s).replace("henry", "henri")
    tokens = re.findall(r"[a-z0-9]+", s)
    while tokens and tokens[0] in {"l", "la", "le", "les", "de", "des", "du", "d"}:
        tokens.pop(0)
    return " ".join(tokens)


def clean(v) -> str | None:
    v = str(v).strip() if v is not None else ""
    return v if v and v.lower() != "none" else None


def ordinal(n: int) -> str:
    return "1ère" if n == 1 else f"{n}ème"


def date_fr(d: date) -> str:
    return f"{JOURS[d.weekday()]} {d.day} {MOIS[d.month - 1]} {d.year}"


def masonic_label(sheet: str) -> str:
    a, b = sheet.split("-")
    return f"{int(a) + 4000}-{int(b) + 4000}"


def meeting_title(n: int, d: date, grade: MeetingGrade, mtype: MeetingType) -> str:
    if n == 0:
        return f"Tenue d'Allumage des Feux du {date_fr(d)}"
    if mtype == MeetingType.FETE:
        return f"{ordinal(n)} Tenue de table du {date_fr(d)}"
    chambre = {
        MeetingGrade.APPRENTI: "en Loge d'App∴",
        MeetingGrade.COMPAGNON: "en Chambre de Comp∴",
        MeetingGrade.MAITRE: "en Chambre du Milieu",
    }[grade]
    suffix = {MeetingType.INITIATION: " — Tenue de réception", MeetingType.PASSAGE: " — Passage"}.get(mtype, "")
    return f"{ordinal(n)} Ten∴ Sol∴ du {date_fr(d)} {chambre}{suffix}"


# ── Lecture des fichiers Excel ───────────────────────────────────────────────

def read_membres(path: str):
    """→ liste des tenues [(sheet, date, grade)] et pointages {(nom, prénom): {(date, grade): code}}."""
    wb = openpyxl.load_workbook(path, data_only=True)
    meetings, records = [], defaultdict(dict)
    for sheet in SHEETS:
        ws = wb[sheet]
        dates, labels, grades = ([c.value for c in ws[i]] for i in (1, 2, 3))
        cols = []
        for i in range(4, len(dates)):
            if not isinstance(dates[i], datetime) or str(labels[i]).strip().upper() != "X":
                continue  # colonne vide ou événement hors tenue (Fresque du Climat)
            d = MEMBRES_DATE_FIXES.get(dates[i].date(), dates[i].date())
            g = GRADE_FROM_EXCEL[ascii_key(grades[i])]
            cols.append((i, d, g))
            meetings.append((sheet, d, g))
        for row in ws.iter_rows(min_row=4, values_only=True):
            name = row[2]
            if not name or str(name).startswith(("Total", "Nb", "%")):
                continue
            who = split_excel_name(name)
            for i, d, g in cols:
                code = clean(row[i])
                if code:
                    records[who][(d, g)] = STATUS_FROM_EXCEL[code.upper()]
    return meetings, records


def read_passants(path: str, meetings_by_date: dict[date, list]):
    """→ fiches visiteurs {clé: infos} et visites [(clé, (date, grade))]."""
    wb = openpyxl.load_workbook(path, data_only=True)
    people, visits, unmatched = {}, [], Counter()
    for sheet in SHEETS:
        ws = wb[sheet]
        months, days = [c.value for c in ws[1]], [c.value for c in ws[2]]
        cols, current, seen = [], None, Counter()
        for i in range(7, len(days)):
            if isinstance(months[i], datetime):
                current = months[i]
            if current is None or not isinstance(days[i], (int, float)):
                continue
            d = date(current.year, current.month, int(days[i]))
            d = PASSANTS_DATE_FIXES.get(d, d)
            # 2 tenues le même jour : même ordre de colonnes que le livre des membres
            rank = seen[d]
            seen[d] += 1
            same_day = meetings_by_date.get(d, [])
            cols.append((i, d, same_day[rank] if rank < len(same_day) else None))
        for row in ws.iter_rows(min_row=3, values_only=True):
            if not row[2]:
                continue
            key = person_key(row[2], row[3])
            people[key] = {  # dernière année connue = informations les plus récentes
                "last_name": clean(row[2]), "first_name": clean(row[3]) or "",
                "grade": clean(row[1]), "lodge": clean(row[4]),
                "obedience": clean(row[5]), "orient": clean(row[6]),
            }
            for i, d, mkey in cols:
                if row[i] in (None, ""):
                    continue
                if mkey is None:
                    unmatched[d] += 1
                else:
                    visits.append((key, mkey))
    return people, visits, unmatched


def is_shouting(s: str) -> bool:
    letters = [c for c in s if c.isalpha()]
    return bool(letters) and sum(c.isupper() for c in letters) / len(letters) > 0.6


def canonical_labels(portal_forms: list[str], excel_forms: list[str]) -> dict[str, str]:
    """Pour chaque clé, la graphie retenue : jamais une graphie en majuscules si une
    autre existe, puis celle déjà normalisée sur le portail, puis la plus fréquente."""
    by_key = defaultdict(Counter)
    for form in portal_forms + excel_forms:
        by_key[label_key(form)][form] += 1
    on_portal = set(portal_forms)
    return {
        k: max(c.items(), key=lambda fc: (not is_shouting(fc[0]), fc[0] in on_portal, fc[1], fc[0]))[0]
        for k, c in by_key.items()
    }


# ── Import ───────────────────────────────────────────────────────────────────

async def run(db: AsyncSession, membres_path: str, passants_path: str) -> None:
    report = Counter()

    # 1. Années maçonniques
    years = {y.label: y for y in (await db.execute(select(MasonicYear))).scalars()}
    year_of_sheet = {}
    for sheet in SHEETS:
        label = masonic_label(sheet)
        if label not in years:
            a, b = (int(x) for x in sheet.split("-"))
            years[label] = MasonicYear(label=label, start_date=date(a, 9, 1), end_date=date(b, 6, 30), is_current=False)
            db.add(years[label])
            report["années créées"] += 1
        year_of_sheet[sheet] = years[label]
    await db.flush()

    # 2. Membres : corrections, membres partis, correspondance des noms
    members = list((await db.execute(select(Member))).scalars())
    for m in members:
        new = RENAMES.get((m.last_name, m.first_name))
        if new:
            print(f"  ✎ Renommé : {m.last_name} {m.first_name} → {new[0]} {new[1]}")
            m.last_name, m.first_name = new
        old_start, new_start = GENERIC_START_FIX
        if m.membership_start_date == old_start:
            m.membership_start_date = new_start
            report["dates d'entrée 22/12 → 21/12/2022"] += 1
    fixes = {person_key(*who): f for who, f in MEMBER_FIXES.items()}
    for m in members:
        for attr, value in fixes.get(person_key(m.last_name, m.first_name), {}).items():
            if getattr(m, attr) != value:
                print(f"  ✎ {attr} {m.last_name} {m.first_name} : {getattr(m, attr)} → {value}")
                setattr(m, attr, value)
    by_key = {person_key(m.last_name, m.first_name): m for m in members}

    excel_meetings, records = read_membres(membres_path)

    first_seen = {who: min(d for d, _ in recs) for who, recs in records.items()}
    for (last, first), (civility, left_on) in DEPARTED.items():
        if person_key(last, first) in by_key:
            continue
        slug = f"{ascii_key(first)}.{ascii_key(last)}".lower()
        m = Member(
            civility=civility, last_name=last, first_name=first,
            email=f"archive.{slug}@amisdesocrate.invalid",  # pas de compte, aucun envoi
            masonic_grade=MasonicGrade.MAITRE, status=MemberStatus.RESIGNED,
            membership_start_date=first_seen.get((last, first)), status_date=left_on,
            program_optin=False, email_notifications=False,
        )
        db.add(m)
        by_key[person_key(last, first)] = m
        print(f"  + Membre parti : {civility}∴ {last} {first} (entré le {m.membership_start_date}, démission le {left_on})")
        report["membres partis créés"] += 1
    await db.flush()

    def resolve(who):
        target = MEMBER_ALIASES.get(who, who)
        return by_key.get(person_key(*target))

    unresolved = [f"{l} {f}" for (l, f) in records if not resolve((l, f))]
    if unresolved:
        raise SystemExit(f"Membres introuvables : {', '.join(unresolved)} — compléter MEMBER_ALIASES/DEPARTED.")

    # 3. Tenues (numérotation chronologique depuis l'allumage des feux = n°0)
    existing = {(m.meeting_date, m.grade): m for m in (await db.execute(select(Meeting))).scalars()}
    initiation_dates = {m.initiation_date for m in members if m.initiation_date}
    passage_dates = {m.companion_date for m in members if m.companion_date}

    plan = [(sheet, d, g, n) for n, (sheet, d, g) in enumerate(sorted(excel_meetings, key=lambda x: x[1]))]
    for d, g, n in EXTRA_MEETINGS:
        sheet = next(s for s, dd, _, _ in plan if dd == d)
        plan.append((sheet, d, g, n))

    meeting_of: dict[tuple, Meeting] = {}
    for sheet, d, g, n in plan:
        if (d, g) in existing:
            meeting_of[(d, g)] = existing[(d, g)]
            report["tenues déjà présentes"] += 1
            continue
        if (d, g) in TABLE_MEETINGS:
            mtype = MeetingType.FETE
        elif g == MeetingGrade.APPRENTI and d in initiation_dates:
            mtype = MeetingType.INITIATION
        elif g == MeetingGrade.COMPAGNON and d in passage_dates:
            mtype = MeetingType.PASSAGE
        else:
            mtype = MeetingType.SOLENNELLE
        mtg = Meeting(
            masonic_year_id=year_of_sheet[sheet].id, meeting_date=d, meeting_number=n,
            title=meeting_title(n, d, g, mtype), grade=g, type=mtype,
            is_locked=True, registration_open=False, agape_enabled=False,
        )
        db.add(mtg)
        meeting_of[(d, g)] = mtg
        report["tenues créées"] += 1
    await db.flush()

    # 4. Pointages des membres
    existing_att = {(a.meeting_id, a.member_id): a for a in (await db.execute(select(Attendance))).scalars()}
    for who, recs in records.items():
        mbr = resolve(who)
        for mkey, status in recs.items():
            mtg = meeting_of[mkey]
            att = existing_att.get((mtg.id, mbr.id))
            if att:
                if att.status != status:
                    print(f"  ⚠ Pointage divergent conservé : {mbr.last_name} {mkey[0]} base={att.status.value} excel={status.value}")
                report["pointages déjà présents"] += 1
                continue
            db.add(Attendance(meeting_id=mtg.id, member_id=mbr.id, status=status))
            report[f"pointages créés ({status.value})"] += 1

    # 5. Maçons passants
    meetings_by_date = defaultdict(list)
    for sheet, d, g in excel_meetings:  # ordre des colonnes du livre des membres
        meetings_by_date[d].append((d, g))
    people, visits, unmatched = read_passants(passants_path, meetings_by_date)
    for d, n in sorted(unmatched.items()):
        print(f"  ⚠ {n} visite(s) le {d} sans tenue correspondante — ignorée(s)")

    visitors = list((await db.execute(select(Visitor))).scalars())
    lodges = canonical_labels([v.lodge_name for v in visitors if v.lodge_name],
                              [p["lodge"] for p in people.values() if p["lodge"]])
    orients = canonical_labels([v.orient_city for v in visitors if v.orient_city],
                               [p["orient"] for p in people.values() if p["orient"]])

    for v in visitors:  # harmonisation des fiches existantes
        for attr, table in (("lodge_name", lodges), ("orient_city", orients)):
            old = getattr(v, attr)
            new = table.get(label_key(old)) if old else None
            if new and new != old:
                print(f"  ✎ {attr} : « {old} » → « {new} »")
                setattr(v, attr, new)
                report["libellés harmonisés"] += 1

    visitor_by_key = {}
    for v in visitors:
        visitor_by_key.setdefault(person_key(v.last_name, v.first_name), v)
    for key, p in people.items():
        grade = VISITOR_GRADE.get(ascii_key(p["grade"])) if p["grade"] else None
        fields = {
            "lodge_name": lodges.get(label_key(p["lodge"])) if p["lodge"] else None,
            "orient_city": orients.get(label_key(p["orient"])) if p["orient"] else None,
            "obedience": p["obedience"].upper() if p["obedience"] else None,
            "masonic_grade": grade,
        }
        v = visitor_by_key.get(key)
        if v:
            for attr, val in fields.items():  # on complète sans écraser
                if val and not getattr(v, attr):
                    setattr(v, attr, val)
            report["passants rattachés à une fiche existante"] += 1
        else:
            v = Visitor(last_name=p["last_name"], first_name=p["first_name"],
                        is_vm=ascii_key(p["grade"]) == "VM", **fields)
            db.add(v)
            visitor_by_key[key] = v
            report["passants créés"] += 1
    await db.flush()

    existing_mv = {(mv.meeting_id, mv.visitor_id) for mv in (await db.execute(select(MeetingVisitor))).scalars()}
    for key, mkey in visits:
        pair = (meeting_of[mkey].id, visitor_by_key[key].id)
        if pair in existing_mv:
            report["visites déjà présentes"] += 1
            continue
        db.add(MeetingVisitor(meeting_id=pair[0], visitor_id=pair[1], status=VisitorStatus.CONFIRMED))
        existing_mv.add(pair)
        report["visites créées"] += 1

    # 6. Rattachement des tracés archivés (GED) aux tenues
    # Colonnes seules : certaines fiches GED anciennes ont un statut hors énumération
    docs = (await db.execute(
        select(Document.id, Document.name, Document.mime_type).where(Document.deleted_at.is_(None))
    )).all()
    tr = re.compile(r"trac[ée] tenue (\d{3})\s*-?\s*(\d{1,2}) (\w+)\.? (\d{4})\s*(\w*)", re.I)
    best: dict[tuple, Document] = {}
    for doc in docs:
        m = tr.search(doc.name or "")
        if not m:
            continue
        n, day, mon, year, suffix = int(m[1]), int(m[2]), ascii_key(m[3]).lower(), int(m[4]), ascii_key(m[5])
        if suffix not in {"", "APP", "COMP", "COM", "MAITRE", "TABLE"}:
            continue  # pièce annexe (ex. « Participants inscrits »)
        month = MOIS_ABR.get(mon) or MOIS_ABR.get(mon[:3])
        if not month:
            print(f"  ⚠ Date illisible dans le nom du tracé : {doc.name}")
            continue
        d = date(year, month, day)
        if d > max(k[0] for k in meeting_of):
            continue  # tracé de l'année en cours, hors périmètre
        g = {"COMP": MeetingGrade.COMPAGNON, "COM": MeetingGrade.COMPAGNON,
             "MAITRE": MeetingGrade.MAITRE}.get(suffix, None)
        candidates = [mk for mk in meeting_of if mk[0] == d and (g is None or mk[1] == g)]
        if g is None and len(candidates) > 1:  # App / Table : la tenue au 1er degré
            candidates = [mk for mk in candidates if mk[1] == MeetingGrade.APPRENTI]
        if len(candidates) != 1:
            print(f"  ⚠ Tracé non rattaché : {doc.name}")
            continue
        mtg = meeting_of[candidates[0]]
        if mtg.meeting_number is not None and mtg.meeting_number != n:
            print(f"  ⚠ Numéro divergent : tracé n°{n} / tenue n°{mtg.meeting_number} ({doc.name})")
        kind = "pdf" if doc.mime_type == "application/pdf" else "word"
        best.setdefault(candidates[0], {}).setdefault(kind, doc)
    reports = {r.meeting_id for r in (await db.execute(select(MeetingReport))).scalars()}
    for mkey, found in best.items():
        mtg = meeting_of[mkey]
        if mtg.id in reports:
            continue
        word = found.get("word")
        if "pdf" in found:
            # Le PDF archivé est la pièce officielle ; le Word, le document de travail
            db.add(MeetingReport(meeting_id=mtg.id, status=ReportStatus.ARCHIVE,
                                 signed_pdf_doc_id=found["pdf"].id,
                                 archived_doc_id=word.id if word else None))
            report["tracés rattachés (PDF officiel)"] += 1
        else:
            db.add(MeetingReport(meeting_id=mtg.id, status=ReportStatus.APPROUVE, archived_doc_id=word.id))
            report["tracés rattachés (Word seul)"] += 1
    missing = sorted(f"n°{m.meeting_number} du {k[0]}" for k, m in meeting_of.items() if k not in best)
    if missing:
        print(f"  ⚠ Tenues sans tracé dans la GED : {', '.join(missing)}")

    await db.flush()
    print("\n── Bilan ──")
    for k, v in report.items():
        print(f"  {k:45s} {v}")


async def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    parser.add_argument("membres", help="Livre des chantiers SRP.xlsx")
    parser.add_argument("passants", help="Livre des chantiers Maçons passants SRP.xlsx")
    parser.add_argument("--apply", action="store_true", help="écrire en base (sinon simulation)")
    args = parser.parse_args()

    url = get_settings().database_url
    print(f"Base : {url.split('@')[-1]}")
    print("Mode : ÉCRITURE" if args.apply else "Mode : SIMULATION (rien n'est écrit)")
    engine = create_async_engine(url)
    try:
        async with async_sessionmaker(engine, expire_on_commit=False)() as db:
            await run(db, args.membres, args.passants)
            if args.apply:
                await db.commit()
                print("\n✅ Import enregistré.")
            else:
                await db.rollback()
                print("\nSimulation terminée — relancer avec --apply pour enregistrer.")
    finally:
        await engine.dispose()


if __name__ == "__main__":
    asyncio.run(main())
